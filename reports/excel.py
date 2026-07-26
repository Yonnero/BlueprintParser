import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from core.models import Blueprint


class ExcelReport:
    def __init__(self, blueprints: list[Blueprint], output_path: str):
        self.blueprints = blueprints
        self.output_path = output_path

    @staticmethod
    def _sort_data(df: pd.DataFrame) -> pd.DataFrame:
        df['global_group'] = df['Тип'].apply(lambda x: 1 if x == "Деталь" else 0)
        type_weight = {"Сборка": 1, "Спецификация": 2, "Деталь": 3}
        df['sort_weight'] = df['Тип'].map(type_weight)

        df.sort_values(by=['global_group', 'Номер', 'sort_weight'], inplace=True)
        df.drop(columns=['global_group', 'sort_weight'], inplace=True)
        df.insert(0, '№', range(1, len(df) + 1))
        return df

    def _apply_styles(self):
        wb = openpyxl.load_workbook(self.output_path)
        ws = wb.worksheets[0]

        bold_font = Font(bold=True)
        link_font = Font(color="0563C1", underline="single")
        center_align = Alignment(horizontal='center', vertical='center')
        left_align = Alignment(horizontal='left', vertical='center')
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                             top=Side(style='thin'), bottom=Side(style='thin'))

        for row in ws.iter_rows(min_row=1, max_row=1):
            for cell in row:
                cell.font = bold_font
                cell.alignment = center_align
                cell.border = thin_border

        pairs_map = {}
        for idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            num = row[1].value
            name = row[2].value
            doc_type = row[3].value

            if num and doc_type in ["Сборка", "Спецификация"]:
                key = (str(num), str(name))
                if key not in pairs_map:
                    pairs_map[key] = {"Сборка": [], "Спецификация": []}
                pairs_map[key][doc_type].append(idx)

        rows_to_highlight = {}
        color_index = 0
        highlight_colors = [
            'FFFFE0', 'E6FFED', 'FFF0E6', 'E6F5FF', 'F9E6FF', 'FFE6F0', 'F0F0F0'
        ]

        for key, types in pairs_map.items():
            if types["Сборка"] and types["Спецификация"]:
                hex_color = highlight_colors[color_index % len(highlight_colors)]
                fill = PatternFill(start_color=hex_color, end_color=hex_color, fill_type='solid')
                for row_idx in types["Сборка"] + types["Спецификация"]:
                    rows_to_highlight[row_idx] = fill
                color_index += 1

        for idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            row[0].alignment = center_align
            row_fill = rows_to_highlight.get(idx)

            for cell in row:
                cell.border = thin_border
                if cell.column != 1:
                    cell.alignment = left_align
                if row_fill:
                    cell.fill = row_fill

            if row[4].value == "-":
                row[4].alignment = center_align

            for col_idx in [5, 6]:
                link_cell = row[col_idx]
                if link_cell.value:
                    link_cell.hyperlink = link_cell.value
                    link_cell.value = "Открыть"
                    link_cell.font = link_font
                    link_cell.alignment = center_align

        ws.column_dimensions['A'].width = 5
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 45
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 40
        ws.column_dimensions['F'].width = 10
        ws.column_dimensions['G'].width = 15
        ws.column_dimensions['H'].width = 30

        wb.save(self.output_path)

    def create(self):
        data = [doc.to_dict() for doc in self.blueprints]
        df = pd.DataFrame(data)
        df = self._sort_data(df)
        df.to_excel(self.output_path, index=False)
        self._apply_styles()